"""Generate CityShield Team Contributions .docx document."""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime


def set_cell_shading(cell, color_hex):
    """Set background color on a table cell."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_bottom_border(paragraph, color="1B3A5C", size="12"):
    """Add a bottom border line under a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_paragraph_spacing(paragraph, before=0, after=0, line=None):
    """Set spacing for a paragraph."""
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = Pt(line)


def create_document():
    doc = Document()

    # -- Page setup --
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    # -- Define styles --
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
    style.paragraph_format.line_spacing = Pt(16)

    # ===========================
    # TITLE PAGE
    # ===========================

    # Top spacer
    for _ in range(6):
        spacer = doc.add_paragraph()
        set_paragraph_spacing(spacer, before=0, after=0)

    # Project name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CityShield")
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)
    run.font.name = "Calibri"
    set_paragraph_spacing(p, before=0, after=4)

    # Subtitle line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Smart City Cyber Range Platform")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
    run.font.italic = True
    set_paragraph_spacing(p, before=0, after=30)

    # Decorative line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("_" * 60)
    run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)
    run.font.size = Pt(11)
    set_paragraph_spacing(p, before=0, after=30)

    # Document title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Team Member Contributions")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)
    set_paragraph_spacing(p, before=0, after=60)

    # Date
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(datetime.datetime.now().strftime("%B %Y"))
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
    set_paragraph_spacing(p, before=0, after=0)

    # Page break
    doc.add_page_break()

    # ===========================
    # TABLE OF CONTENTS (simple)
    # ===========================

    p = doc.add_paragraph()
    run = p.add_run("Table of Contents")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)
    set_paragraph_spacing(p, before=0, after=6)
    add_bottom_border(p)

    toc_items = [
        ("1.", "Introduction"),
        ("2.", "Team Overview"),
        ("3.", "Individual Contributions"),
        ("", "3.1  Sami — Team Leader"),
        ("", "3.2  Bati & Khaled — Researchers & Penetration Testing"),
        ("", "3.3  Thamer — Security Analyst"),
        ("", "3.4  Abdulwahab — UI/UX & Visualization"),
        ("4.", "Collaboration & Workflow"),
        ("5.", "Conclusion"),
    ]

    for num, title in toc_items:
        p = doc.add_paragraph()
        if num:
            run = p.add_run(f"{num}  {title}")
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
        else:
            run = p.add_run(f"      {title}")
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
        set_paragraph_spacing(p, before=2, after=2, line=18)

    doc.add_page_break()

    # ===========================
    # HELPER: Section heading
    # ===========================

    def add_section_heading(text, level=1):
        p = doc.add_paragraph()
        run = p.add_run(text)
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x5C)
            set_paragraph_spacing(p, before=6, after=6)
            add_bottom_border(p)
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2A, 0x52, 0x7A)
            set_paragraph_spacing(p, before=12, after=4)
        elif level == 3:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x3A, 0x6A, 0x9A)
            set_paragraph_spacing(p, before=8, after=4)
        return p

    def add_body_text(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
        set_paragraph_spacing(p, before=0, after=6, line=18)
        p.paragraph_format.first_line_indent = Cm(0)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style="List Bullet")
        if bold_prefix:
            run = p.add_run(bold_prefix)
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
            run = p.add_run(text)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
        else:
            run = p.add_run(text)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
        set_paragraph_spacing(p, before=1, after=1, line=17)
        return p

    # ===========================
    # 1. INTRODUCTION
    # ===========================

    add_section_heading("1.  Introduction")

    add_body_text(
        "This document outlines the individual and collaborative contributions of each team member "
        "to the CityShield project. CityShield is a smart city cyber range platform developed for "
        "training, testing, and evaluating cybersecurity defences within simulated urban infrastructure "
        "environments. The platform integrates traffic management simulation, IoT sensor networks, "
        "network emulation, threat detection aligned to the MITRE ATT&CK framework, and automated "
        "incident response capabilities."
    )

    add_body_text(
        "The purpose of this document is to provide a transparent and detailed record of each member's "
        "responsibilities, technical contributions, and areas of ownership within the project. It is "
        "intended for academic assessment and serves to demonstrate the division of labour, individual "
        "expertise, and the collaborative effort that brought the platform to completion."
    )

    # ===========================
    # 2. TEAM OVERVIEW
    # ===========================

    add_section_heading("2.  Team Overview")

    add_body_text(
        "The CityShield development team consisted of five members, each assigned a primary role "
        "based on their skills and interests. The following table summarises the team composition:"
    )

    # Team overview table
    table = doc.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    table.autofit = True

    # Header row
    headers = ["Team Member", "Primary Role", "Key Focus Areas"]
    header_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        header_cells[i].text = ""
        p = header_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(header_cells[i], "1B3A5C")

    data = [
        ("Sami", "Team Leader", "Project management, IoT, network, security awareness"),
        ("Bati & Khaled", "Researchers & Pen Testers", "Attack research, penetration testing, MITRE mapping"),
        ("Thamer", "Security Analyst", "Alerts, detection rules, analyst workflow"),
        ("Abdulwahab", "UI/UX & Visualization", "Interface design, 3D smart city map, user experience"),
    ]

    for row_idx, (name, role, focus) in enumerate(data, start=1):
        cells = table.rows[row_idx].cells
        for col_idx, val in enumerate([name, role, focus]):
            cells[col_idx].text = ""
            p = cells[col_idx].paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(10)
            if col_idx == 0:
                run.font.bold = True
            if row_idx % 2 == 0:
                set_cell_shading(cells[col_idx], "EDF2F7")

    # Set column widths
    for row in table.rows:
        row.cells[0].width = Cm(3.5)
        row.cells[1].width = Cm(4.5)
        row.cells[2].width = Cm(9.0)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=0)

    # ===========================
    # 3. INDIVIDUAL CONTRIBUTIONS
    # ===========================

    add_section_heading("3.  Individual Contributions")

    add_body_text(
        "The following subsections detail the specific contributions made by each team member, "
        "including the components they developed, the technical challenges they addressed, "
        "and their impact on the overall system."
    )

    # --- 3.1 Sami ---
    add_section_heading("3.1  Sami — Team Leader", level=2)

    add_body_text(
        "As the Team Leader, Sami was responsible for overseeing the project's overall direction, "
        "coordinating between team members, and ensuring that all components integrated cohesively "
        "into a unified platform. Beyond the leadership role, Sami made direct technical contributions "
        "across multiple subsystems, filling gaps wherever additional development effort was required."
    )

    add_section_heading("Leadership & Project Management", level=3)
    add_bullet("Defined the project scope, milestones, and deliverable timeline")
    add_bullet("Coordinated task assignments and tracked progress across all workstreams")
    add_bullet("Facilitated team communication and resolved technical blockers")
    add_bullet("Ensured consistent code quality, documentation standards, and integration testing")

    add_section_heading("Technical Contributions", level=3)
    add_bullet(
        "Developed the IoT device integration layer, including sensor simulation "
        "and data ingestion into the OpenSearch data store"
    )
    add_bullet(
        "Contributed to the network architecture design, defining the three-network Docker topology "
        "(main bridge, isolated cyber range, isolated IoT range) that underpins service isolation"
    )
    add_bullet(
        "Built the Security Awareness training module — an interactive portal with training content "
        "covering phishing, password security, data protection, and incident reporting, along with "
        "a knowledge assessment quiz"
    )
    add_bullet(
        "Assisted in implementing select attack scenarios, ensuring they generated authentic traffic "
        "patterns suitable for detection rule evaluation"
    )
    add_bullet(
        "Addressed cross-cutting integration issues across frontend, backend, and microservice layers "
        "to maintain system coherence"
    )

    # --- 3.2 Bati & Khaled ---
    add_section_heading("3.2  Bati & Khaled — Researchers & Penetration Testing", level=2)

    add_body_text(
        "Bati and Khaled jointly owned the offensive security research and attack simulation "
        "components of the platform. Their work provided the realistic threat scenarios that form "
        "the core training value of CityShield."
    )

    add_section_heading("Cybersecurity Research", level=3)
    add_bullet(
        "Conducted in-depth research on real-world attack techniques targeting smart city "
        "infrastructure, including traffic management systems, IoT sensor networks, and industrial "
        "control systems"
    )
    add_bullet(
        "Analysed threat intelligence sources and academic literature to identify the most relevant "
        "attack vectors for the platform's training scenarios"
    )

    add_section_heading("Penetration Testing", level=3)
    add_bullet(
        "Performed penetration testing against the platform's simulated infrastructure to validate "
        "vulnerability exposure and detection coverage"
    )
    add_bullet(
        "Tested the cyber range environment (Metasploitable target) and IoT range (sensor hub) "
        "to ensure realistic exploitability for training purposes"
    )

    add_section_heading("Attack Scenario Design & Implementation", level=3)
    add_bullet(
        "Designed multi-stage attack scenarios spanning the full cyber kill chain — from initial "
        "reconnaissance and enumeration through to lateral movement, data exfiltration, and impact"
    )
    add_bullet(
        "Built the attack stages for each scenario, ensuring progressive difficulty and "
        "realistic attacker behaviour patterns"
    )
    add_bullet(
        "Mapped all attack scenarios and detection rules to the MITRE ATT&CK framework, "
        "covering 31 techniques across multiple tactics (Reconnaissance, Credential Access, "
        "Lateral Movement, Exfiltration, Impact, and others)"
    )
    add_bullet(
        "Contributed to the Custom Scenario Builder's technique catalog, defining per-technique "
        "parameters for researcher-configurable attack chains"
    )

    # --- 3.3 Thamer ---
    add_section_heading("3.3  Thamer — Security Analyst", level=2)

    add_body_text(
        "Thamer focused on the defensive side of the platform, developing the analyst-facing "
        "workflows that enable security professionals to detect, investigate, and respond to "
        "threats within the CityShield environment."
    )

    add_section_heading("Analyst Workflow Development", level=3)
    add_bullet(
        "Designed and implemented the analyst investigation workflow, including the alert "
        "triage process, status management (open, acknowledged, resolved, closed), and "
        "the tabbed investigation interface (Analysis, Actions, Related Events)"
    )
    add_bullet(
        "Developed the alert analysis pipeline, integrating MITRE ATT&CK context, threat "
        "intelligence enrichment data, and remediation recommendations into each alert view"
    )

    add_section_heading("Threat Detection & Response", level=3)
    add_bullet(
        "Contributed to the implementation of detection rules, working on rule logic, "
        "threshold tuning, and false positive documentation"
    )
    add_bullet(
        "Supported the development of the alert response logic, including the manual action "
        "execution workflow and the confirmation safeguards"
    )
    add_bullet(
        "Worked on alert severity classification and the event replay timeline feature "
        "that allows analysts to review the events surrounding an alert trigger"
    )

    # --- 3.4 Abdulwahab ---
    add_section_heading("3.4  Abdulwahab — UI/UX & Visualization", level=2)

    add_body_text(
        "Abdulwahab was responsible for the platform's visual identity and user experience, "
        "creating an interface that makes complex cybersecurity data accessible and engaging."
    )

    add_section_heading("User Interface Design", level=3)
    add_bullet(
        "Designed the overall UI architecture, establishing the cyberpunk-themed visual identity "
        "with glassmorphism effects, neon accent colours, and dark backgrounds"
    )
    add_bullet(
        "Developed the responsive layout system, navigation sidebar with role-based menu items, "
        "and page transition animations using Framer Motion"
    )
    add_bullet(
        "Created the dashboard components, including statistics cards, charts (Recharts), "
        "and the data table interfaces used across Alerts, Rules, Devices, and Scenarios pages"
    )

    add_section_heading("3D Smart City Visualization", level=3)
    add_bullet(
        "Developed the interactive 3D smart city map using React Three Fiber and Three.js, "
        "featuring procedurally generated buildings grouped into six zones (Traffic, IoT, Network, "
        "Security, Industrial, Cyber Range)"
    )
    add_bullet(
        "Implemented data-driven building visuals where building height reflects risk score "
        "and colour reflects operational status (green for healthy, amber for warning, red for "
        "critical, grey for offline)"
    )
    add_bullet(
        "Built the real-time attack visualisation system — when a scenario runs, targeted "
        "buildings flash with pulsing effects, point lights, and rotating ground rings to "
        "provide immediate visual feedback"
    )
    add_bullet(
        "Implemented the asset inspector panel (click-to-inspect), floating markers, and "
        "camera controls for an intuitive exploration experience"
    )

    add_section_heading("User Experience Enhancements", level=3)
    add_bullet(
        "Improved overall usability through consistent interaction patterns, loading states, "
        "confirmation dialogs, and error handling across all pages"
    )
    add_bullet(
        "Optimised 3D rendering performance through material instancing, selective post-processing, "
        "and efficient per-frame update patterns"
    )

    # ===========================
    # 4. COLLABORATION & WORKFLOW
    # ===========================

    add_section_heading("4.  Collaboration & Workflow")

    add_body_text(
        "The team adopted a collaborative development approach with clear ownership boundaries. "
        "Each member was responsible for their designated components while maintaining open "
        "communication channels for cross-functional integration. Key collaborative practices "
        "included:"
    )

    add_bullet(
        "Version control via Git and GitHub, with feature branches and pull requests for code review"
    )
    add_bullet(
        "Regular team meetings to synchronise progress, identify dependencies, and resolve blockers"
    )
    add_bullet(
        "Shared documentation standards (CLAUDE.md, README.md, docs/ directory) to ensure "
        "knowledge transfer and onboarding efficiency"
    )
    add_bullet(
        "Integration testing across component boundaries — attack scenarios (Bati & Khaled) "
        "validated against detection rules (Thamer), visualised in the 3D city (Abdulwahab), "
        "and orchestrated through the platform pipeline (Sami)"
    )
    add_bullet(
        "Continuous integration via GitHub Actions ensuring code quality "
        "(linting, testing, build validation) on every commit"
    )

    # ===========================
    # 5. CONCLUSION
    # ===========================

    add_section_heading("5.  Conclusion")

    add_body_text(
        "The CityShield platform is the result of a structured and collaborative team effort, "
        "with each member contributing specialised expertise to deliver a comprehensive smart city "
        "cyber range. Sami provided overall leadership and cross-system integration; Bati and Khaled "
        "built the offensive security layer with realistic attack scenarios mapped to MITRE ATT&CK; "
        "Thamer developed the defensive analyst workflows and detection capabilities; and Abdulwahab "
        "created the visual interface and 3D city that brings the platform to life."
    )

    add_body_text(
        "The clear division of responsibilities, combined with consistent communication and shared "
        "tooling, enabled the team to deliver a platform that spans the full cybersecurity operations "
        "lifecycle — from attack simulation through detection to automated response — within an "
        "engaging and interactive training environment."
    )

    # ===========================
    # SAVE
    # ===========================

    output_path = "/Users/sami/Documents/GitHub/CityShield/CityShield_Team_Contributions.docx"
    doc.save(output_path)
    print(f"Document saved to: {output_path}")


if __name__ == "__main__":
    create_document()
